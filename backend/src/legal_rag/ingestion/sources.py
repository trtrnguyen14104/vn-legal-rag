import hashlib
import io
import json
import re
import time
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse
from urllib.robotparser import RobotFileParser
from zoneinfo import ZoneInfo

import httpx
from bs4 import BeautifulSoup

from legal_rag.config import settings
from legal_rag.domain.legal import DATE, NUMBER, OFFICIAL_HOSTS, date_strings, parse_date
from legal_rag.domain.models import DocumentRecord
from legal_rag.domain.text import normalize

DOWNLOAD_HOSTS = OFFICIAL_HOSTS | {"datafiles.chinhphu.vn", "congbaocdn.chinhphu.vn", "g7.cdnchinhphu.vn"}


class SourceError(RuntimeError):
    pass


def record_review_errors(record):
    errors = []
    if record.source_kind not in ("legislation", "consolidated"):
        errors.append("unverified_source_kind")
    if record.scope != "central" or re.search(r"UBND|HĐND", record.number, re.I):
        errors.append("outside_central_scope")
    url = urlparse(record.source_url)
    if (
        url.scheme != "https"
        or url.hostname not in OFFICIAL_HOSTS | {"vbpl-bientap-gateway.moj.gov.vn"}
        or url.username
    ):
        errors.append("missing_official_provenance")
    if not record.license_evidence:
        errors.append("missing_license_evidence")
    if not re.search(r"(?im)^\s*Điều\s+\d", record.text):
        errors.append("missing_legal_structure")
    for field in ("effective_from", "effective_to"):
        value = getattr(record, field)
        if value:
            try:
                parsed = parse_date(value)
                proof = getattr(record, field + "_evidence")
                dates = [parse_date(d) for d in date_strings(proof or "")]
                if not proof or parsed not in dates:
                    errors.append(field + "_without_evidence")
            except ValueError:
                errors.append(field + "_invalid")
    try:
        start, end = parse_date(record.effective_from), parse_date(record.effective_to)
        if start and end and end <= start:
            errors.append("invalid_effective_interval")
    except ValueError:
        pass
    return errors


def extract_binary(content, suffix):
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        if len(reader.pages) > 1000:
            raise SourceError("pdf_too_many_pages")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if len(text.strip()) < 100:
            raise SourceError("ocr_required")
        return text
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        return (
            "\n".join(p.text for p in document.paragraphs)
            + "\n"
            + "\n".join(
                " | ".join(c.text for c in row.cells) for table in document.tables for row in table.rows
            )
        )
    return content.decode("utf-8-sig")


def read_records(path):
    """Explicit canonical adapter; unknown dataset schemas are quarantined, never guessed."""
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        with path.open(encoding="utf-8-sig") as stream:
            for line in stream:
                if line.strip():
                    yield json.loads(line)
    elif suffix == ".json":
        raw = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(raw, list):
            yield from raw
        elif isinstance(raw, dict) and "documents" in raw:
            yield from raw["documents"]
        elif isinstance(raw, dict):
            yield raw
        else:
            raise SourceError("unknown_json_schema")
    elif suffix == ".parquet":
        import pyarrow.parquet as pq

        for batch in pq.ParquetFile(path).iter_batches(batch_size=256):
            yield from batch.to_pylist()
    elif suffix in (".pdf", ".docx", ".txt", ".html"):
        sidecar = path.with_suffix(path.suffix + ".metadata.json")
        if not sidecar.exists():
            raise SourceError("metadata_sidecar_required")
        raw = json.loads(sidecar.read_text(encoding="utf-8"))
        text = extract_binary(path.read_bytes(), suffix)
        if suffix == ".html":
            soup = BeautifulSoup(text, "html.parser")
            for node in soup(["script", "style", "nav", "footer"]):
                node.decompose()
            text = soup.get_text("\n", strip=True)
        yield raw | {"text": text}
    else:
        raise SourceError("unsupported_format")


class OfficialConnector:
    def __init__(self, source):
        self.source = source
        self.host = "vanban.chinhphu.vn" if source == "vanban" else "congbao.chinhphu.vn"
        self.robots = {}
        self.failures = []

    def fetch(self, url):
        config = settings()
        for _ in range(4):
            parsed = urlparse(url)
            if (
                parsed.scheme != "https"
                or parsed.hostname not in DOWNLOAD_HOSTS
                or parsed.username
                or parsed.port not in (None, 443)
            ):
                raise SourceError("unsafe_source_url")
            base = "https://" + parsed.hostname
            if base not in self.robots:
                with httpx.Client(timeout=30, follow_redirects=False) as client:
                    response = client.get(
                        base + "/robots.txt", headers={"User-Agent": config.source_user_agent}
                    )
                if response.status_code not in (200, 404):
                    raise SourceError("robots_unavailable")
                robot = RobotFileParser()
                robot.parse(
                    response.text.splitlines()
                    if response.status_code == 200
                    else ["User-agent: *", "Allow: /"]
                )
                self.robots[base] = robot
            if not self.robots[base].can_fetch(config.source_user_agent, url):
                raise SourceError("robots_disallowed")
            time.sleep(config.source_delay_seconds)
            for attempt in range(3):
                with httpx.Client(timeout=30, follow_redirects=False) as client:
                    with client.stream(
                        "GET", url, headers={"User-Agent": config.source_user_agent}
                    ) as response:
                        if response.status_code in (429, 502, 503, 504) and attempt < 2:
                            time.sleep(2**attempt)
                            continue
                        if response.is_redirect:
                            url = urljoin(url, response.headers["location"])
                            break
                        response.raise_for_status()
                        pieces, size = [], 0
                        for piece in response.iter_bytes():
                            size += len(piece)
                            if size > config.source_max_bytes:
                                raise SourceError("source_too_large")
                            pieces.append(piece)
                        return b"".join(pieces)
            else:
                raise SourceError("source_retry_exhausted")
        raise SourceError("redirect_limit")

    def is_detail(self, url):
        parsed = urlparse(url)
        if parsed.hostname != self.host:
            return False
        if self.source == "vanban":
            return "docid" in parse_qs(parsed.query)
        return bool(re.search(r"/(?:van-ban|vb)/.+\d", parsed.path)) or "docid" in parse_qs(parsed.query)

    def listings(self, since):
        config = settings()
        pending = [u for u in config.source_list_urls.split(",") if urlparse(u).hostname == self.host]
        seen, details = set(), set()
        while pending:
            if len(seen) >= config.source_max_pages:
                self.failures.append("listing_page_limit")
                break
            url = pending.pop(0)
            if url in seen:
                continue
            seen.add(url)
            soup = BeautifulSoup(self.fetch(url), "html.parser")
            for anchor in soup.select("a[href]"):
                link = urljoin(url, anchor["href"])
                if self.is_detail(link):
                    container = anchor.find_parent("tr") or anchor.parent
                    dates = DATE.findall(container.get_text(" ", strip=True))
                    issued = None
                    if dates:
                        try:
                            issued = parse_date(dates[0])
                        except ValueError:
                            pass
                    if issued is None or issued >= since:
                        details.add(link)
                elif urlparse(link).hostname == self.host and (
                    anchor.get("rel") == ["next"]
                    or re.search(r"(?:trang sau|tiếp|next)", anchor.get_text(" ", strip=True), re.I)
                ):
                    pending.append(link)
            # Some sites use postback pagination; mark incomplete rather than claiming coverage.
            if soup.select("a[href*='__doPostBack']"):
                self.failures.append("postback_pagination_requires_adapter")
        if not details:
            raise SourceError("no_document_links_or_schema_changed")
        return sorted(details)

    def detail(self, url):
        content = self.fetch(url)
        soup = BeautifulSoup(content, "html.parser")
        labels = {}
        for row in soup.select("tr"):
            cells = row.find_all(["td", "th"], recursive=False)
            if len(cells) == 2:
                labels[normalize(cells[0].get_text(" ", strip=True)).rstrip(":").lower()] = normalize(
                    cells[1].get_text(" ", strip=True)
                )
        title_node = soup.select_one("h1")
        title = title_node.get_text(" ", strip=True) if title_node else ""
        number = labels.get("số ký hiệu") or labels.get("số hiệu")
        if not number:
            matches = NUMBER.findall(title)
            number = matches[0] if matches else ""
        attachments = []
        for node in soup.select("a[href],a[data-href],iframe[src],embed[src]"):
            link = urljoin(url, node.get("data-href") or node.get("href") or node.get("src", ""))
            filename = (
                node.get("data-file")
                or parse_qs(urlparse(link).query).get("file_name", [urlparse(link).path])[0]
            )
            suffix = Path(filename).suffix.lower()
            if not suffix and filename.lower().endswith("pdf"):
                suffix = ".pdf"
            if urlparse(link).hostname in DOWNLOAD_HOSTS and suffix in (".pdf", ".docx"):
                attachments.append((link, suffix))
        texts, artifacts = [], [(url, content, ".html")]
        for attachment, suffix in dict.fromkeys(attachments):
            binary = self.fetch(attachment)
            texts.append(extract_binary(binary, suffix))
            artifacts.append((attachment, binary, suffix))
        if not texts:
            article = soup.select_one(".fulltext, .toanvan, #toanvan, .doc-content, .document-content")
            if article:
                texts = [article.get_text("\n", strip=True)]
        if not title or not number or not texts:
            raise SourceError("missing_document_title_number_or_body")
        kind = labels.get("loại văn bản") or next(
            (
                k
                for k in ("Nghị định", "Thông tư", "Luật", "Quyết định", "Nghị quyết", "Văn bản hợp nhất")
                if k.lower() in title.lower()
            ),
            "unknown",
        )
        effective_from = labels.get("ngày có hiệu lực") or labels.get("ngày hiệu lực")
        effective_to = labels.get("ngày hết hiệu lực")
        raw = {
            "title": title[:300],
            "number": number,
            "kind": kind,
            "text": "\n".join(texts),
            "source_url": url,
            "authority": labels.get("cơ quan ban hành", ""),
            "source_kind": "consolidated" if "VBHN" in number.upper() else "legislation",
            "effective_from": effective_from or None,
            "effective_to": effective_to or None,
            "effective_from_evidence": "Ngày có hiệu lực: " + effective_from if effective_from else None,
            "effective_to_evidence": "Ngày hết hiệu lực: " + effective_to if effective_to else None,
            "issued_at": labels.get("ngày ban hành"),
            "license_evidence": f"Nguồn công bố chính thức: {url}; ghi rõ nguồn khi tái sử dụng.",
            "verified_at": datetime.now(ZoneInfo("UTC")),
        }
        return DocumentRecord.model_validate(raw), artifacts


def bucket_manifest(output):
    from huggingface_hub import HfApi

    config = settings()
    api = HfApi(token=config.hf_token.get_secret_value() or False)
    report = {"checked_at": datetime.now(ZoneInfo("UTC")).isoformat(), "buckets": []}
    for bucket in config.hf_buckets.split(","):
        try:
            info = api.bucket_info(bucket)
            entries = list(api.list_bucket_tree(bucket, recursive=True))
            files = [
                {"path": e.path, "size": e.size, "type": e.type, "hash": getattr(e, "xet_hash", None)}
                for e in entries
            ]
            report["buckets"].append(
                {"id": bucket, "size": info.size, "files": files, "status": "accessible"}
            )
        except Exception as exc:
            report["buckets"].append(
                {"id": bucket, "status": "unavailable", "error_code": type(exc).__name__}
            )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def sha256_file(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def update_window(checkpoint):
    today = datetime.now(ZoneInfo("Asia/Saigon")).date()
    if checkpoint and checkpoint.get("checked_at"):
        return checkpoint["checked_at"].date() - timedelta(days=7)
    return today - timedelta(days=7)
